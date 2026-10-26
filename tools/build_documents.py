from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 49, 63)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color in (("Title", 24, "123B61"), ("Heading 1", 16, "123B61"), ("Heading 2", 12.5, "2467A6"), ("Heading 3", 11, "2467A6")):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("代码哨兵 CodeGuard  |  项目材料")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(120, 136, 153)


def title_page(doc: Document, subtitle: str):
    p = doc.add_paragraph()
    p.style = "Title"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("代码哨兵 CodeGuard")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(36, 103, 166)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.add_run("强网杯大学生创新创业作品赛道").bold = True
    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    entries = [("项目名称", "代码哨兵 CodeGuard"), ("作品方向", "代码敏感信息泄露风险审计"), ("团队信息", "待填：团队成员、学校、指导教师"), ("文档用途", "比赛申报与项目展示")]
    for row, (key, value) in zip(meta.rows, entries):
        row.cells[0].text = key
        row.cells[1].text = value
        set_cell_shading(row.cells[0], "EAF2F8")
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    doc.add_page_break()


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, headers):
        cell.text = text
        set_cell_shading(cell, "123B61")
        set_cell_margins(cell, 120, 110, 120, 110)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
    for row_data in rows:
        row = table.add_row()
        for cell, text in zip(row.cells, row_data):
            cell.text = str(text)
            set_cell_margins(cell, 100, 110, 100, 110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph()
    return table


def build_proposal():
    doc = Document()
    configure(doc)
    title_page(doc, "项目申报书")
    add_heading(doc, "项目概述", 1)
    add_para(doc, "代码哨兵 CodeGuard 是一个面向大学生和教师的代码敏感信息泄露风险审计平台。用户可以粘贴代码或上传 ZIP 项目，平台通过原创的混合检测策略识别 API Key、Access Key / Secret Key、数据库账号密码、JWT Token、邮箱、手机号和疑似高熵凭据，自动计算风险等级并定位到文件和代码行。平台同时提供人工复核、中文修复建议、正确示例代码和 PDF 审计报告，帮助学生在提交代码前建立密钥保护和隐私保护意识。")
    add_para(doc, "本项目的核心边界是辅助发现和教学解释。系统不会自动修改源代码；调用 DeepSeek 时只向接口发送临时脱敏上下文，原始代码留在本地。")

    add_heading(doc, "一 项目背景", 1)
    add_para(doc, "在课程作业、实验项目和毕业设计中，学生经常把 API Key、数据库密码、令牌或个人信息直接写进源代码。问题通常不是因为缺少编程能力，而是因为密钥管理知识没有融入日常编码流程。一旦代码被上传到公开仓库、提交到教学平台或进入截图和文档，凭据就可能被复制、滥用或长期留存在版本历史中。")
    add_para(doc, "现有学习过程往往停留在发现问题之后的口头提醒，学生看到了警告，却不一定理解风险来源和正确改法。CodeGuard 将扫描、解释、人工确认和报告归档放在同一个流程中，适合课程实验、代码作业检查和毕业设计留档。")

    add_heading(doc, "二 用户痛点", 1)
    add_table(doc, ["痛点", "具体表现", "项目回应"], [
        ("发现不及时", "密钥通常在提交或展示之后才被发现", "支持粘贴代码和 ZIP 项目，提交前即可检查"),
        ("定位不直观", "学生只知道存在风险，不知道是哪一行", "展示文件名、行号、风险代码和风险类型"),
        ("建议不易理解", "只提示不要泄露，没有可学习的替代写法", "给出中文解释和正确示例代码"),
        ("整改难留痕", "课程作业或毕设缺少可归档的安全记录", "支持人工状态标记和 PDF 审计报告"),
    ], widths=[1.2, 2.3, 2.7])

    add_heading(doc, "三 产品功能", 1)
    add_bullets(doc, [
        "多入口输入：支持文本粘贴和 ZIP 压缩包上传，覆盖单文件和多文件项目。",
        "多语言识别：首版覆盖 Python、Java、C++、JavaScript，并读取常见配置文件。",
        "分级检测：将结果划分为严重、高、中、低四级，并显示风险评分和置信度。",
        "位置高亮：按文件和行号展示风险代码，便于课堂演示和人工复核。",
        "人工复核：支持待确认、已确认、误报、已修复、忽略本次五种状态。",
        "修复教学：提供文字版修复建议和正确示例代码，可选调用 DeepSeek 生成更具体的解释。",
        "报告归档：一键生成包含风险概览、问题详情和整改记录的 PDF 审计报告。",
    ])

    add_heading(doc, "四 技术路线", 1)
    add_para(doc, "系统采用 Python Flask 后端和静态网页前端。核心检测在本地完成，DeepSeek API 作为可选的建议增强模块。")
    add_table(doc, ["模块", "实现方式", "输出"], [
        ("输入层", "文本粘贴与 ZIP 安全解压，过滤二进制和常见依赖目录", "待扫描的代码文件集合"),
        ("检测层", "规则信号、变量上下文、文件类型、占位符判断和高熵特征组合", "风险线索、行号、证据"),
        ("评分层", "按风险类别设置基础分，并叠加上下文权重和占位符扣分", "严重、高、中、低风险"),
        ("解释层", "内置中文建议，或把临时脱敏上下文发送到 DeepSeek", "修复建议和正确示例"),
        ("展示层", "静态 HTML、CSS、JavaScript 页面，支持风险选择和状态更新", "代码高亮与整改进度"),
        ("报告层", "ReportLab 生成中文 PDF，自动汇总扫描统计和整改记录", "审计报告文件"),
    ], widths=[1.2, 3.4, 1.6])

    add_heading(doc, "五 创新点", 1)
    add_para(doc, "1. 检测结果与教学解释合并。项目不仅提示风险，还通过上下文证据、风险评分、修复建议和示例代码解释为什么危险以及怎样改。")
    add_para(doc, "2. 面向学生的人工复核闭环。系统将误报、已修复和忽略本次纳入界面状态，使扫描结果能够转化为可追踪的学习记录。")
    add_para(doc, "3. 规则检测与大模型建议分层。核心检测不依赖外部模型，接口不可用时仍能完成扫描；大模型只增强解释，不拥有修改源代码的权限。")
    add_para(doc, "4. 将敏感信息审计转化为可归档成果。PDF 报告适用于课程作业、毕业设计和安全教育活动，方便教师批改和学生自查。")
    add_para(doc, "5. 原创评分与规则组织。检测器以项目自定义的风险类别、上下文权重和占位符识别为基础，不照搬外部项目的规则表和核心检测流程。")

    add_heading(doc, "六 市场与应用价值", 1)
    add_para(doc, "项目首先服务高校课程、实验室和毕业设计过程，也可作为网络安全社团、编程教学平台和校园创新实践活动的辅助工具。对于教师，平台降低了批量检查作业中常见凭据问题的成本；对于学生，平台把抽象的密钥安全要求变成可观察、可解释、可复核的代码问题。")
    add_para(doc, "后续可以扩展自定义检测规则、教学班级统计、代码仓库流水线接入、密钥轮换提醒以及更多配置格式支持。项目暂不承诺验证密钥是否真实有效，也不替代企业级密钥管理系统。")

    add_heading(doc, "七 实施计划", 1)
    add_table(doc, ["阶段", "主要工作", "阶段成果"], [
        ("第一阶段", "完成输入、检测、分级和结果展示", "可运行的本地网页 MVP"),
        ("第二阶段", "完成人工复核、DeepSeek 建议和 PDF 报告", "完整演示流程"),
        ("第三阶段", "补充测试、演示数据、用户手册和比赛材料", "可提交项目包"),
        ("后续迭代", "扩展规则、语言和课堂管理能力", "面向教学场景的版本升级"),
    ], widths=[1.2, 3.6, 1.4])

    add_heading(doc, "八 风险分析", 1)
    add_table(doc, ["风险", "影响", "应对方式"], [
        ("误报或漏报", "影响学生判断", "保留置信度和人工状态，不把结果表述为绝对结论"),
        ("真实密钥外传", "造成二次泄露", "发送给 DeepSeek 前仅使用临时脱敏上下文，原始代码不出本地"),
        ("网络或接口不可用", "建议增强功能暂时不可用", "本地检测和内置建议独立工作"),
        ("规则覆盖有限", "无法识别所有凭据类型", "支持后续添加团队自定义规则和样例测试"),
        ("材料误含个人信息", "不符合盲审脱敏要求", "申报书、PPT 和展示材料使用占位符并进行提交前检查"),
    ], widths=[1.4, 2.1, 2.7])

    add_heading(doc, "原创性与知识产权说明", 1)
    add_para(doc, "本项目的页面、风险评分方式、规则组织、修复建议模板、演示数据和材料文本均为本团队围绕高校代码安全教育场景重新设计。项目不直接复制他人论文、开源项目 README、源代码、规则清单或核心检测逻辑。使用第三方依赖时，仅用于通用的 Web 服务、PDF 生成和测试，不替代本项目的检测实现。")
    add_heading(doc, "演示流程", 1)
    add_bullets(doc, ["粘贴包含 API Key、数据库密码和个人信息的演示代码。", "查看风险总览和四级分级结果，点击问题跳转到对应文件和代码行。", "将问题标记为已确认或误报，展示整改闭环。", "打开修复建议，说明环境变量和受控配置的正确写法。", "导出 PDF 审计报告，说明可用于课程作业和毕设存档。"])
    path = OUT / "CodeGuard_项目申报书.docx"
    doc.save(path)
    return path


def build_manual():
    doc = Document()
    configure(doc)
    title_page(doc, "用户手册")
    add_heading(doc, "一 产品说明", 1)
    add_para(doc, "CodeGuard 代码安全审计平台帮助学生和教师在代码提交前发现敏感信息泄露线索。平台支持本地网页访问，扫描器在本机完成，DeepSeek 只负责可选的建议优化。")
    add_heading(doc, "二 安装环境", 1)
    add_bullets(doc, ["Windows 10 或 Windows 11。", "Python 3.10 及以上版本。", "允许联网时可以配置 DeepSeek API；未配置时仍可使用本地检测。", "浏览器使用 Chrome、Edge 或其他现代浏览器。"])
    add_heading(doc, "三 启动系统", 1)
    add_para(doc, "在项目根目录打开 PowerShell，依次执行以下命令：")
    p = doc.add_paragraph(style="Intense Quote")
    p.add_run("python -m venv .venv\n.\\.venv\\Scripts\\Activate.ps1\npip install -r requirements.txt\nCopy-Item .env.example .env\npython run.py")
    add_para(doc, "启动后在浏览器打开 http://127.0.0.1:5000。")
    add_heading(doc, "四 粘贴代码检测", 1)
    add_bullets(doc, ["选择“粘贴代码”标签。", "填写代码文件名，例如 demo.py。", "粘贴待检测代码，点击“开始检测”。", "系统展示风险数量、安全评分、文件位置和问题清单。"])
    add_heading(doc, "五 ZIP 项目检测", 1)
    add_bullets(doc, ["选择“上传 ZIP”标签并选择项目压缩包。", "单个 ZIP 最大 20 MB，最多读取 500 个文件。", "系统会忽略 .git、node_modules、虚拟环境、构建目录和二进制文件。", "压缩包路径会经过安全校验，避免读取路径穿越目标。"])
    add_heading(doc, "六 阅读和确认风险", 1)
    add_para(doc, "点击右侧问题清单中的一条风险，左侧会显示文件名、行号、风险类型和原始代码行。列表中的置信度用于辅助判断，不代表密钥一定有效。可以把问题标记为已确认、误报、已修复或忽略本次。")
    add_heading(doc, "七 查看修复建议", 1)
    add_para(doc, "系统默认提供内置中文建议和正确示例。配置 DeepSeek 后，可以点击“使用 DeepSeek 优化建议”。发送给模型的只是临时脱敏上下文，系统不会把真实敏感值作为建议请求发送，也不会自动修改本地源码。")
    add_heading(doc, "八 配置 DeepSeek", 1)
    add_para(doc, "复制 .env.example 为 .env，填写 DEEPSEEK_API_KEY。可以按需要修改 DEEPSEEK_MODEL 和 DEEPSEEK_BASE_URL。请不要把真实 Key 提交到代码仓库、截图、PDF 或申报材料中。")
    add_heading(doc, "九 导出 PDF 报告", 1)
    add_para(doc, "完成扫描后点击“导出 PDF 报告”。报告包含扫描时间、文件和代码行统计、安全评分、风险等级分布、问题详情、修复建议和整改记录。建议在完成人工确认后导出最终版。")
    add_heading(doc, "十 使用边界", 1)
    add_bullets(doc, ["系统是教育性代码安全辅助工具，不替代正式安全审计。", "系统不验证密钥是否有效，也不执行任何外部系统操作。", "系统不会自动修改或覆盖源代码。", "发现真实密钥后，应立即轮换凭据，并检查版本历史和日志。"])
    add_heading(doc, "十一 常见问题", 1)
    add_table(doc, ["问题", "处理方式"], [
        ("页面无法打开", "确认 run.py 仍在运行，并访问 http://127.0.0.1:5000。"),
        ("没有发现风险", "检查文件后缀是否支持，并进行人工复核；检测器不保证覆盖所有凭据格式。"),
        ("DeepSeek 建议失败", "检查网络和 .env 配置；本地检测与内置建议不受影响。"),
        ("ZIP 无法扫描", "确认文件是 ZIP，大小不超过 20 MB，并且包含支持的文本代码文件。"),
    ], widths=[2.1, 4.1])
    path = OUT / "CodeGuard_用户手册.docx"
    doc.save(path)
    return path


def build_test_report():
    doc = Document()
    configure(doc)
    title_page(doc, "测试报告")
    add_heading(doc, "一 测试目的", 1)
    add_para(doc, "验证 CodeGuard 的核心扫描逻辑、风险分级、ZIP 输入、人工状态更新、PDF 报告接口和 DeepSeek 未配置时的降级行为，确认系统可以支撑比赛现场演示和课程作业使用。")
    add_heading(doc, "二 测试环境", 1)
    add_table(doc, ["项目", "配置"], [("操作系统", "Windows"), ("Python", "3.12"), ("后端", "Flask"), ("报告", "ReportLab"), ("测试框架", "pytest"), ("模型接口", "DeepSeek，可选")], widths=[1.6, 4.6])
    add_heading(doc, "三 自动化测试结果", 1)
    add_table(doc, ["编号", "测试内容", "预期结果", "结果"], [
        ("T01", "识别多个敏感类别", "API Key、数据库密码、邮箱、手机号和 JWT 被识别", "通过"),
        ("T02", "占位符降权", "your_key_here 等样例值不直接判为高风险", "通过"),
        ("T03", "忽略依赖目录", "node_modules 等目录不进入扫描结果", "通过"),
        ("T04", "安全评分计算", "存在高风险时评分低于 100", "通过"),
        ("T05", "文本接口冒烟", "接口返回扫描结果和风险摘要", "通过"),
        ("T06", "ZIP 接口冒烟", "ZIP 内可支持文件被扫描", "通过"),
        ("T07", "PDF 接口冒烟", "接口返回 application/pdf", "通过"),
    ], widths=[0.6, 2.0, 2.8, 0.8])
    add_heading(doc, "四 关键行为验证", 1)
    add_bullets(doc, ["系统在本地完成基础扫描，不需要 DeepSeek 才能输出结果。", "未配置 API Key 时，风险项仍然可以查看内置中文建议和示例代码。", "调用 DeepSeek 的代码只使用 finding.masked_line 作为临时上下文，未实现自动修改或回写接口。", "ZIP 解压过程拒绝绝对路径和包含 .. 的路径，限制单文件大小和文件数量。", "PDF 报告默认展示脱敏代码行，降低报告二次传播风险。"])
    add_heading(doc, "五 演示数据", 1)
    add_para(doc, "data/demo_project.zip 为教学演示数据，包含虚构的 API Key、数据库密码、JWT Token、邮箱、手机号和云访问密钥。所有值仅用于演示，不对应真实账号或真实服务。")
    add_heading(doc, "六 已知限制", 1)
    add_bullets(doc, ["规则检测不能覆盖所有自定义凭据格式，存在误报和漏报可能。", "当前结果保存在进程内存，重启服务后不会保留历史扫描记录。", "DeepSeek 建议依赖外部网络和用户自己的 API 配额。", "系统暂未接入真实 Git 平台、企业密钥库或持续集成流水线。"])
    add_heading(doc, "七 结论", 1)
    add_para(doc, "当前版本已经覆盖比赛演示所需主流程：输入代码、发现风险、定位代码行、查看建议、人工确认和导出 PDF。自动化测试与接口冒烟测试均通过，可以在继续补充现场演示脚本和材料脱敏检查后用于项目展示。")
    path = OUT / "CodeGuard_测试报告.docx"
    doc.save(path)
    return path


def build_core_info():
    doc = Document()
    configure(doc)
    title_page(doc, "项目核心信息")
    add_heading(doc, "项目名称", 1)
    add_para(doc, "代码哨兵 CodeGuard")
    add_heading(doc, "一句话简介", 1)
    add_para(doc, "面向大学生和教师的代码敏感信息泄露审计平台，帮助用户在提交代码前发现密钥、令牌和个人信息风险，并通过中文解释和示例代码完成安全学习。")
    add_heading(doc, "项目简介", 1)
    add_para(doc, "CodeGuard 支持文本粘贴和 ZIP 项目上传，覆盖 Python、Java、C++、JavaScript。系统以本地混合检测为核心，结合风险类别、上下文和高熵特征进行评分，按严重、高、中、低四级展示结果。用户可以查看文件和行号、标记人工处理状态、获得修复建议并生成 PDF 审计报告。")
    add_heading(doc, "创新点", 1)
    add_bullets(doc, ["将风险发现、代码定位、修复解释和审计留档串成一个教学闭环。", "采用规则、上下文和高熵特征的组合评分，减少只靠单一正则表达式的误报。", "大模型作为可选解释层，基础检测不依赖模型，且只接收临时脱敏上下文。", "人工复核状态让学生能够记录误报、确认和修复过程。"])
    add_heading(doc, "技术路线", 1)
    add_para(doc, "输入层接收文本或 ZIP，安全解压后过滤路径和无关目录；检测层遍历支持文件并提取候选风险；评分层根据风险类别和上下文计算分数；展示层提供代码行高亮、风险清单和状态更新；解释层提供内置建议和 DeepSeek 建议；报告层生成中文 PDF。")
    add_heading(doc, "五分钟演示口径", 1)
    add_table(doc, ["时间", "展示内容", "表达重点"], [
        ("0:00-0:40", "介绍问题", "学生容易把密钥和个人信息写进代码"),
        ("0:40-1:40", "粘贴演示代码", "一次检测发现多类风险并定位行号"),
        ("1:40-2:40", "查看分级与建议", "风险评分、风险原因、环境变量示例"),
        ("2:40-3:40", "人工确认", "误报、已确认、已修复等状态形成闭环"),
        ("3:40-4:30", "导出 PDF", "课程作业和毕业设计可以留档"),
        ("4:30-5:00", "总结创新", "本地检测稳定，大模型只增强解释，原始源码不自动修改"),
    ], widths=[1.1, 2.2, 3.1])
    path = OUT / "CodeGuard_项目核心信息.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    paths = [build_proposal(), build_manual(), build_test_report(), build_core_info()]
    for path in paths:
        print(path)

